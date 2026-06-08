"""Generate a teaching slide deck (presentation.pptx) covering the local LM
fine-tuning + RAG pipeline in this workspace.

Run:  python scripts/make_ppt.py
Output: presentation.pptx in the repo root.
"""
import os

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# --- Theme ---------------------------------------------------------------
DARK = RGBColor(0x1F, 0x2A, 0x44)
ACCENT = RGBColor(0x2E, 0x6F, 0xF2)
LIGHT = RGBColor(0xF5, 0xF7, 0xFB)
GREY = RGBColor(0x55, 0x5B, 0x6E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def _bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def _box(slide, left, top, width, height):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    return tf


def title_slide(title, subtitle):
    s = prs.slides.add_slide(BLANK)
    _bg(s, DARK)
    tf = _box(s, 0.9, 2.4, 11.5, 2.0)
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = title
    r.font.size = Pt(44); r.font.bold = True; r.font.color.rgb = WHITE
    p2 = tf.add_paragraph()
    r2 = p2.add_run(); r2.text = subtitle
    r2.font.size = Pt(22); r2.font.color.rgb = RGBColor(0xB9, 0xC6, 0xE8)
    return s


def bullet_slide(title, bullets):
    """bullets: list of (text, level) tuples."""
    s = prs.slides.add_slide(BLANK)
    _bg(s, LIGHT)
    # title bar
    bar = s.shapes.add_shape(1, Inches(0), Inches(0), prs.slide_width, Inches(1.15))
    bar.fill.solid(); bar.fill.fore_color.rgb = DARK; bar.line.fill.background()
    ttf = bar.text_frame; ttf.word_wrap = True
    ttf.margin_left = Inches(0.6)
    tp = ttf.paragraphs[0]
    tr = tp.add_run(); tr.text = title
    tr.font.size = Pt(28); tr.font.bold = True; tr.font.color.rgb = WHITE

    tf = _box(s, 0.8, 1.5, 11.7, 5.6)
    first = True
    for text, level in bullets:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = level
        r = p.add_run(); r.text = text
        if level == 0:
            r.font.size = Pt(20); r.font.bold = True; r.font.color.rgb = DARK
        else:
            r.font.size = Pt(16); r.font.color.rgb = GREY
        p.space_after = Pt(6)
    return s


def two_col_slide(title, left_head, left_items, right_head, right_items):
    s = prs.slides.add_slide(BLANK)
    _bg(s, LIGHT)
    bar = s.shapes.add_shape(1, Inches(0), Inches(0), prs.slide_width, Inches(1.15))
    bar.fill.solid(); bar.fill.fore_color.rgb = DARK; bar.line.fill.background()
    ttf = bar.text_frame; ttf.margin_left = Inches(0.6)
    tr = ttf.paragraphs[0].add_run(); tr.text = title
    tr.font.size = Pt(28); tr.font.bold = True; tr.font.color.rgb = WHITE

    def col(left, head, items, head_color):
        card = s.shapes.add_shape(1, Inches(left), Inches(1.5), Inches(5.9), Inches(5.4))
        card.fill.solid(); card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = head_color; card.line.width = Pt(1.5)
        tf = card.text_frame; tf.word_wrap = True
        tf.margin_left = Inches(0.3); tf.margin_right = Inches(0.3); tf.margin_top = Inches(0.25)
        hp = tf.paragraphs[0]; hr = hp.add_run(); hr.text = head
        hr.font.size = Pt(22); hr.font.bold = True; hr.font.color.rgb = head_color
        for it in items:
            p = tf.add_paragraph(); r = p.add_run(); r.text = "- " + it
            r.font.size = Pt(15); r.font.color.rgb = GREY; p.space_after = Pt(5)

    col(0.7, left_head, left_items, ACCENT)
    col(6.75, right_head, right_items, RGBColor(0xC0, 0x39, 0x2B))
    return s


def code_slide(title, lines):
    s = prs.slides.add_slide(BLANK)
    _bg(s, LIGHT)
    bar = s.shapes.add_shape(1, Inches(0), Inches(0), prs.slide_width, Inches(1.15))
    bar.fill.solid(); bar.fill.fore_color.rgb = DARK; bar.line.fill.background()
    ttf = bar.text_frame; ttf.margin_left = Inches(0.6)
    tr = ttf.paragraphs[0].add_run(); tr.text = title
    tr.font.size = Pt(28); tr.font.bold = True; tr.font.color.rgb = WHITE

    card = s.shapes.add_shape(1, Inches(0.7), Inches(1.5), Inches(11.9), Inches(5.4))
    card.fill.solid(); card.fill.fore_color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    card.line.fill.background()
    tf = card.text_frame; tf.word_wrap = True
    tf.margin_left = Inches(0.35); tf.margin_top = Inches(0.25)
    first = True
    for ln in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        r = p.add_run(); r.text = ln
        r.font.name = 'Consolas'; r.font.size = Pt(15)
        r.font.color.rgb = RGBColor(0x9C, 0xDC, 0xFE) if ln.strip().startswith('#') else RGBColor(0xE6, 0xED, 0xF3)
        p.space_after = Pt(3)
    return s


def build():
    title_slide(
        "From Fine-Tuning to RAG",
        "Building and understanding a local language-model pipeline (CPU) - Pandya history demo",
    )

    bullet_slide("Agenda", [
        ("1. The data pipeline: documents -> chunks -> training data", 0),
        ("2. Tokenization: turning text into numbers", 0),
        ("3. The model: logits, probabilities, attention", 0),
        ("4. Fine-tuning a small causal LM (distilgpt2)", 0),
        ("5. Generation and why it repeats / hallucinates", 0),
        ("6. RAG: grounding answers in your own documents", 0),
        ("7. Fine-tuning vs RAG, and a commands cheat-sheet", 0),
    ])

    bullet_slide("The Data Pipeline", [
        ("docs/*.txt  ->  chunks.jsonl  ->  train.jsonl  ->  out_model", 0),
        ("prepare_data.py: split each document into paragraph chunks", 1),
        ("Each chunk gets a SHA-1 id so identical text has the same fingerprint", 1),
        ("prepare_lm_data.py: build self-supervised training text", 1),
        ("The label for each token is simply the next token - the text supervises itself", 1),
        ("Same chunks.jsonl is later reused as the RAG knowledge base", 1),
    ])

    bullet_slide("Tokenization - the bridge to the model", [
        ("The tokenizer converts text <-> integer IDs", 0),
        ("GPT-2 uses subword (byte-pair) tokens, so words can split into pieces", 1),
        ("The leading marker means 'a space precedes this token'", 1),
        ("Example: 'cooking pasta.' -> ['cooking', 'pasta', '.']", 1),
        ("Vocabulary size is ~50,257 tokens", 1),
    ])

    bullet_slide("The Model: logits -> probabilities", [
        ("For every position the model outputs a score (logit) for EVERY vocab token", 0),
        ("softmax turns ~50k logits into a probability distribution that sums to 1", 1),
        ("The model's whole job: 'given the text so far, what comes next?'", 1),
        ("Greedy decoding = always pick the single highest-probability token", 1),
    ])

    bullet_slide("Attention - which tokens look at which", [
        ("Each Transformer layer lets a token attend to earlier tokens", 0),
        ("distilgpt2: 6 layers; gpt2/this demo shows 12 heads per layer", 1),
        ("The attention matrix is lower-triangular (causal)", 0),
        ("A causal LM can only look backward - never at future tokens", 1),
        ("That is what makes it generate left-to-right", 1),
    ])

    bullet_slide("Fine-Tuning a Small Causal LM", [
        ("train_lm.py fine-tunes distilgpt2/gpt2 on your chunks (CPU-friendly)", 0),
        ("Key hyperparameters:", 0),
        ("per_device_train_batch_size x gradient_accumulation_steps = effective batch", 1),
        ("max_length = token block size per training example", 1),
        ("num_train_epochs = full passes over the data", 1),
        ("--debug_run caps training at 5 steps (smoke test only, not real training)", 1),
        ("A fine-tuned model = big pretrained brain + small custom nudge", 0),
    ])

    bullet_slide("Generation: greedy vs sampling", [
        ("Generation = predict-next-token in a loop, feeding output back in", 0),
        ("Greedy (default): deterministic, but loops/repeats", 1),
        ("Sampling (--do_sample, temperature): more varied, less looping", 1),
        ("repetition_penalty and no_repeat_ngram_size further reduce loops", 1),
        ("Always pass attention_mask to avoid the pad/eos warning", 1),
    ])

    bullet_slide("Why outputs repeat and hallucinate", [
        ("Repetition: greedy decoding keeps picking the same most-likely token", 0),
        ("Hallucination: the model generates statistical patterns, not facts", 0),
        ("Example: prompt 'Chandravamsha' produced 'Rajiv Gandhi', 'Jitender Singh', dates", 1),
        ("None of those are in your text - they come from distilgpt2's web pretraining", 1),
        ("The model is OFFLINE: no internet lookups, only frozen weights", 0),
    ])

    bullet_slide("Tiny fine-tune: what to expect", [
        ("~23 short chunks + few epochs barely move the weights", 0),
        ("Pretrained web knowledge dominates the output", 1),
        ("Before training: 'pandya rulers' -> off-topic political text", 1),
        ("After 3-epoch training: 'pandya rulers' -> kings/rulers/country/capital flavor", 1),
        ("It learns the TOPIC/STYLE, not reliable FACTS", 0),
    ])

    bullet_slide("RAG - Retrieval-Augmented Generation", [
        ("Keep knowledge OUTSIDE the model, in your documents", 0),
        ("At query time: retrieve relevant chunks and put them in the prompt", 0),
        ("build_embeddings.py: embed chunks (all-MiniLM-L6-v2) -> embeddings.npy", 1),
        ("query_rag.py: embed the question, cosine-search top-k, build a grounded prompt", 1),
        ("NumPy search by default; FAISS optional (no Python 3.14 wheels)", 1),
        ("No retraining needed - update documents anytime", 0),
    ])

    bullet_slide("How RAG reuses the earlier pipeline", [
        ("chunks.jsonl (from prepare_data.py) becomes the retrieval corpus", 0),
        ("out_model (from train_lm.py) is pluggable as the RAG reader", 0),
        ("Generation logic mirrors generate.py: tokenize -> generate -> decode", 0),
        ("RAG only ADDS an embed -> retrieve -> prompt layer in front", 0),
        ("Fine-tuning becomes OPTIONAL: RAG's power is retrieval", 1),
    ])

    bullet_slide("The reader matters: causal vs instruction-tuned", [
        ("Retrieval can be perfect, yet the answer is still garbage", 0),
        ("distilgpt2/out_model is a base completion model - cannot follow instructions", 1),
        ("It ignores context and free-associates from pretrained memory", 1),
        ("google/flan-t5-base is instruction-tuned (seq2seq) and answers FROM context", 0),
        ("query_rag.py auto-detects seq2seq vs causal and decodes accordingly", 1),
        ("--extractive returns the top chunk verbatim: always faithful, no hallucination", 1),
    ])

    two_col_slide(
        "Fine-Tuning vs RAG",
        "Fine-Tuning", [
            "Bakes knowledge INTO the weights",
            "Needs retraining to update",
            "Data-hungry; small data hallucinates",
            "No sources / not citable",
            "Good for STYLE and behavior",
        ],
        "RAG", [
            "Keeps knowledge in documents",
            "Update docs anytime, no retraining",
            "Works with little data",
            "Can cite the retrieved chunks",
            "Good for FACTUAL, grounded answers",
        ],
    )

    code_slide("Commands Cheat-Sheet", [
        "# 1) Build chunks and training data",
        "python scripts/prepare_data.py",
        "python scripts/prepare_lm_data.py",
        "",
        "# 2) Fine-tune (real run)",
        "python scripts/train_lm.py --model_name_or_path distilgpt2 \\",
        "  --num_train_epochs 3 --max_length 256 --save_steps 100",
        "",
        "# 3) Inspect tokens / logits / attention",
        "python scripts/inspect_model.py",
        "",
        "# 4) Build the RAG index",
        "python scripts/build_embeddings.py",
        "",
        "# 5) Ask, grounded in your docs (flan-t5 default)",
        "python scripts/query_rag.py --query \"What is Chandravamsha?\" --top_k 4 --show_context",
        "",
        "# Faithful, no generation:",
        "python scripts/query_rag.py --query \"What is Chandravamsha?\" --extractive --show_context",
    ])

    bullet_slide("Key Takeaways", [
        ("A model only 'knows' what is frozen in its weights - unless you give it context", 0),
        ("Small fine-tunes steer style, not facts, and hallucinate confidently", 0),
        ("RAG grounds answers in your documents and reuses the same pipeline plumbing", 0),
        ("RAG quality needs BOTH a good retriever AND a capable, instruction-tuned reader", 0),
        ("Everything here runs locally on CPU and offline after the one-time downloads", 0),
    ])

    repo_root = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
    out_path = os.path.join(repo_root, 'presentation.pptx')
    prs.save(out_path)
    print('Wrote', out_path, '-', len(prs.slides._sldIdLst), 'slides')


if __name__ == '__main__':
    build()
